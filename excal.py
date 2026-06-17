    
import openpyxl as xl

from openpyxl.chart import BarChart,Reference

def automate_file(filename):
    
    file = xl.load_workbook(filename)
    sheet = file['Stock Summary']

    for row in range(3, sheet.max_row + 1):
        sheet.cell(row,4).value = f"=B{row}*C{row}"
        
    value = Reference(sheet,min_row=4,max_row=sheet.max_row,min_col=4,max_col=4)
    chart = BarChart()
    chart.add_data(value)
    sheet.add_chart(chart,'e3')

    file.save(filename)
    
    
automate_file = ('STOCK_18%(1).xlsx')